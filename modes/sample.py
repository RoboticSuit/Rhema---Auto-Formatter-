# Source Trace:
# File: sample_v0.5.py
# Knowledge Files: CodeSourceDB v3.6, SyntaxBiasDB v2.3, HumanSyntaxDB v1.2
# REF_IDs: RM_DO178_001, RM_NIST_001, RM_HCI_002, WEB_PY_001, WEB_PY_002,
#          WEB_PY_004
# Logic: Generates a ready-to-style format sample .docx and offers to save it
#        to the template library. User may skip the save at any time.
#        v0.5 -- Section breaks added between cover, TOC, body, and references.
#                REFERENCES heading and marker added as section 4.
#                writer.py body scan fixed (contents_found flag).

"""
Rhema -- Auto Formatter
modes/sample.py

Single responsibility: Generates a format sample document demonstrating all
structural elements Rhema supports (headings, body text, table, bullet list,
numbered list). The user opens the sample in Word, applies their own visual
styling, saves it, and feeds it back to Rhema as a template via the library.

After writing the file Rhema offers to add it to the template library
immediately. The user may skip this and add it later via Manage Templates.

Governing standards:
  NIST RMF Step 6 HITL gate     -- REF_ID: RM_NIST_001
  DO-178C state traceability    -- REF_ID: RM_DO178_001
  ISO 9241-110:2020 interaction -- REF_ID: RM_HCI_002
"""

from pathlib import Path
from typing import Any, Optional

from config.constants import (
    DEFAULT_OUTPUT_DIR,
    OUTPUT_EXTENSION,
    SAMPLE_OUTPUT_FILENAME,
    PREVIEW_VALID_CONFIRM_SIGNALS,
)
from storage.template_store import save_template
from ui.messages import (
    MSG_SAMPLE_INTRO,
    MSG_SAMPLE_GENERATING,
    MSG_SAMPLE_SAVE_PROMPT,
    MSG_SAMPLE_SKIPPED,
    MSG_SAMPLE_WRITE_ERROR,
    msg_sample_done,
    MSG_SAVE_TEMPLATE_NAME_PROMPT,
    msg_template_saved,
    fmt,
    LABEL_ERROR,
)


# ---------------------------------------------------------------------------
# Sample document content declarations
# All text is declared as named constants -- no magic strings in logic.
# REF_ID: RM_DO178_001
# ---------------------------------------------------------------------------

_SAMPLE_HEADING_1A: str  = "1. Section Heading"
_SAMPLE_BODY_1A: str     = (
    "This is a body paragraph under a Heading 1. Replace this text with your "
    "own content. Rhema will apply your template styles to paragraphs like this."
)
_SAMPLE_HEADING_2A: str  = "1.1 Subsection Heading"
_SAMPLE_BODY_2A: str     = (
    "This is a body paragraph under a Heading 2. Your template can define "
    "different font sizes, weights, and spacing for each heading level."
)
_SAMPLE_HEADING_1B: str  = "2. Tables"
_SAMPLE_BODY_1B: str     = (
    "The table below demonstrates a standard three-column layout with a "
    "header row. Style the header row and body rows differently in Word."
)
_SAMPLE_TABLE_HEADERS: tuple = ("Column A", "Column B", "Column C")
_SAMPLE_TABLE_ROWS: tuple = (
    ("Row 1 Cell 1", "Row 1 Cell 2", "Row 1 Cell 3"),
    ("Row 2 Cell 1", "Row 2 Cell 2", "Row 2 Cell 3"),
    ("Row 3 Cell 1", "Row 3 Cell 2", "Row 3 Cell 3"),
)
_SAMPLE_HEADING_1C: str  = "3. Lists"
_SAMPLE_HEADING_2B: str  = "3.1 Bullet List"
_SAMPLE_BULLET_ITEMS: tuple = (
    "First bullet point",
    "Second bullet point",
    "Third bullet point",
)
_SAMPLE_HEADING_2C: str  = "3.2 Numbered List"
_SAMPLE_NUMBERED_ITEMS: tuple = (
    "First numbered item",
    "Second numbered item",
    "Third numbered item",
)

# Word style names used in the sample -- default styles present in all .docx files
_STYLE_HEADING_1: str = "Heading 1"
_STYLE_HEADING_2: str = "Heading 2"
_STYLE_NORMAL:    str = "Normal"
_STYLE_BULLET:    str = "List Bullet"
_STYLE_NUMBER:    str = "List Number"

# Cover page table -- label + placeholder pairs.
# Structural placeholders (BODY_CONTENT_START, REFERENCES_START) are excluded
# because they are not cover page fields. REF_ID: RM_DO178_001
_COVER_PAGE_TITLE: str = "COVER PAGE"
_COVER_FIELDS: tuple = (
    ("Document Type",          "[[DOCUMENT_TYPE]]"),
    ("Document Title",         "[[DOCUMENT_TITLE]]"),
    ("Subtitle",               "[[SUBTITLE]]"),
    ("Author Name",            "[[AUTHOR_NAME]]"),
    ("Department / Organisation", "[[DEPARTMENT_ORGANISATION]]"),
    ("Month / Year",           "[[MONTH_YEAR]]"),
    ("Organisation",           "[[ORGANISATION]]"),
    ("Confidentiality Label",  "[[CONFIDENTIALITY_LABEL]]"),
    ("Institution Name",       "[[INSTITUTION_NAME]]"),
    ("Faculty / Department",   "[[FACULTY_DEPARTMENT]]"),
    ("Student ID",             "[[STUDENT_ID]]"),
    ("Supervisor Name",        "[[SUPERVISOR_NAME]]"),
    ("Course Code",            "[[COURSE_CODE]]"),
    ("Course Name",            "[[COURSE_NAME]]"),
    ("Submission Date",        "[[SUBMISSION_DATE]]"),
)

# Column header labels for the cover page table
_COVER_TABLE_HEADER_FIELD: str = "Field"
_COVER_TABLE_HEADER_PH:    str = "Placeholder (edit or delete rows as needed)"

# Structural marker paragraphs -- required by formatter.py as injection points.
# These must be present in every template. Do not remove them.
# REF_ID: RM_DO178_001
_MARKER_BODY_START:  str = "[[BODY_CONTENT_START]]"
_MARKER_BODY_NOTE:   str = "*** Do not remove the line above. Rhema injects your document body here. ***"
_MARKER_REFS_START:  str = "[[REFERENCES_START]]"
_MARKER_REFS_NOTE:   str = "*** Do not remove the line above. Rhema injects your references here. ***"

# TOC section constants
_TOC_HEADING_TEXT:   str = "CONTENTS"
_TOC_INSTR_TEXT:     str = " TOC \\h \\o \"1-3\" "   # identical to Fix A in writer.py
_STYLE_TITLE:        str = "Title"

# References section constants
_REFS_HEADING_TEXT:  str = "REFERENCES"


# ---------------------------------------------------------------------------
# Cover page builder
# REF_ID: RM_DO178_001, WEB_PY_001
# ---------------------------------------------------------------------------

def _build_cover_page(doc: Any) -> None:
    """
    Insert a cover page section as the first content in the sample document.

    Adds a Heading 1 title followed by a two-column table. Each row contains
    a plain-text field label in the left cell and the corresponding
    [[PLACEHOLDER]] token in the right cell.

    The user can restyle the table, add rows with custom placeholders, or
    delete rows for fields they do not need. Rhema reads whatever placeholders
    are present in the saved template and prompts only for those.
    REF_ID: RM_DO178_001, RM_NIST_001
    """
    doc.add_paragraph(_COVER_PAGE_TITLE, style=_STYLE_HEADING_1)

    # Two-column table: Field label | Placeholder token
    # Row 0 is the column header row.
    row_count: int = len(_COVER_FIELDS) + 1   # header row + one row per field
    cover_table = doc.add_table(rows=row_count, cols=2)

    # Header row
    header_cells = cover_table.rows[0].cells
    header_cells[0].text = _COVER_TABLE_HEADER_FIELD
    header_cells[1].text = _COVER_TABLE_HEADER_PH

    # Field rows -- one per declared cover placeholder
    for row_idx, (label, placeholder) in enumerate(_COVER_FIELDS):
        cells = cover_table.rows[row_idx + 1].cells
        cells[0].text = label
        cells[1].text = placeholder

    # Spacer paragraph between cover page and body content
    doc.add_paragraph()


# ---------------------------------------------------------------------------
# TOC section builder
# REF_ID: RM_DO178_001, WEB_PY_001
# ---------------------------------------------------------------------------

def _build_toc_section(doc: Any) -> None:
    """
    Insert a CONTENTS heading and a TOC field paragraph into the document.

    The CONTENTS paragraph uses the Title style so _post_process_docx in
    writer.py recognises and excludes it from the TOC itself.

    The TOC field uses the same instruction established by Fix A:
      TOC \\h \\o "1-3"
    \\h -- entries are hyperlinks (Ctrl+Click navigation)
    \\o "1-3" -- includes outline levels 1 through 3

    The field is marked dirty=true so Word updates it automatically on open
    or when the COM session calls toc.Update(). REF_ID: RM_DO178_001
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    # CONTENTS heading paragraph -- Title style
    doc.add_paragraph(_TOC_HEADING_TEXT, style=_STYLE_TITLE)

    # TOC field paragraph built from raw OOXML.
    # Structure: fldChar(begin, dirty) | instrText | fldChar(separate) | placeholder | fldChar(end)
    # REF_ID: RM_DO178_001, WEB_PY_001
    toc_para = doc.add_paragraph()
    toc_para.style = doc.styles[_STYLE_NORMAL]
    p = toc_para._p

    def _make_run(parent: Any) -> Any:
        r = OxmlElement("w:r")
        parent.append(r)
        return r

    # Run 1: fldChar begin (dirty=true so Word rebuilds on open)
    r1 = _make_run(p)
    fc_begin = OxmlElement("w:fldChar")
    fc_begin.set(qn("w:fldCharType"), "begin")
    fc_begin.set(qn("w:dirty"),       "true")
    r1.append(fc_begin)

    # Run 2: instrText containing the TOC instruction
    r2 = _make_run(p)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = _TOC_INSTR_TEXT
    r2.append(instr)

    # Run 3: fldChar separate
    r3 = _make_run(p)
    fc_sep = OxmlElement("w:fldChar")
    fc_sep.set(qn("w:fldCharType"), "separate")
    r3.append(fc_sep)

    # Run 4: placeholder text (replaced by Word on update)
    r4 = _make_run(p)
    t = OxmlElement("w:t")
    t.text = " "
    r4.append(t)

    # Run 5: fldChar end
    r5 = _make_run(p)
    fc_end = OxmlElement("w:fldChar")
    fc_end.set(qn("w:fldCharType"), "end")
    r5.append(fc_end)


# ---------------------------------------------------------------------------
# Section break helper
# Inserts an inline sectPr (next-page section break) into the last paragraph
# of the current section. This gives each section its own page in the output.
# REF_ID: RM_DO178_001, WEB_PY_001
# ---------------------------------------------------------------------------

def _add_section_break(doc: Any) -> None:
    """
    Append a next-page section break after the current last paragraph.

    In OOXML, an inline sectPr inside a paragraph's pPr creates a section
    boundary. The paragraph itself becomes the last paragraph of the section
    and the next paragraph starts a new page section.
    REF_ID: RM_DO178_001
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    # Add a blank paragraph to carry the section break
    break_para = doc.add_paragraph()
    pPr = OxmlElement("w:pPr")
    sectPr = OxmlElement("w:sectPr")
    pgSz = OxmlElement("w:pgSz")
    # A4 page size in twentieths of a point (OOXML units)
    pgSz.set(qn("w:w"),      "11906")   # 210mm
    pgSz.set(qn("w:h"),      "16838")   # 297mm
    pgSz.set(qn("w:orient"), "portrait")
    sectPr.append(pgSz)
    pPr.append(sectPr)
    break_para._p.insert(0, pPr)


# ---------------------------------------------------------------------------
# Sample document builder
# REF_ID: RM_DO178_001, WEB_PY_001
# ---------------------------------------------------------------------------

def _build_sample_doc() -> Optional[Any]:
    """
    Build and return a python-docx Document containing all structural elements.
    Starts with a cover page table, then headings, body, content table, and lists.
    Returns None if python-docx is not available.
    REF_ID: RM_DO178_001
    """
    try:
        from docx import Document as DocxDocument
    except ImportError:
        return None

    doc = DocxDocument()

    # Remove the blank paragraph python-docx adds by default
    if len(doc.paragraphs) > 0:
        p = doc.paragraphs[0]._element
        p.getparent().remove(p)

    # --- Cover page table section (Section 1) ---
    _build_cover_page(doc)
    _add_section_break(doc)   # Section 1 ends here -- next page starts TOC

    # --- TOC section (Section 2) ---
    # CONTENTS heading + TOC field. COM session finds and updates this.
    # REF_ID: RM_DO178_001
    _build_toc_section(doc)
    _add_section_break(doc)   # Section 2 ends here -- next page starts body

    # --- Body content injection marker (Section 3) ---
    # This paragraph is required. formatter.py replaces it with the document's
    # body content. Style it as Normal -- it will not appear in the output.
    # REF_ID: RM_DO178_001
    doc.add_paragraph(_MARKER_BODY_START, style=_STYLE_NORMAL)
    doc.add_paragraph(_MARKER_BODY_NOTE,  style=_STYLE_NORMAL)

    # --- Heading 1A and body ---
    doc.add_paragraph(_SAMPLE_HEADING_1A, style=_STYLE_HEADING_1)
    doc.add_paragraph(_SAMPLE_BODY_1A,    style=_STYLE_NORMAL)

    # --- Heading 2A and body ---
    doc.add_paragraph(_SAMPLE_HEADING_2A, style=_STYLE_HEADING_2)
    doc.add_paragraph(_SAMPLE_BODY_2A,    style=_STYLE_NORMAL)

    # --- Heading 1B (Tables section) and intro body ---
    doc.add_paragraph(_SAMPLE_HEADING_1B, style=_STYLE_HEADING_1)
    doc.add_paragraph(_SAMPLE_BODY_1B,    style=_STYLE_NORMAL)

    # --- Table ---
    col_count: int = len(_SAMPLE_TABLE_HEADERS)
    row_count: int = len(_SAMPLE_TABLE_ROWS) + 1   # header row + data rows
    table = doc.add_table(rows=row_count, cols=col_count)

    # Header row
    header_cells = table.rows[0].cells
    for col_idx, header_text in enumerate(_SAMPLE_TABLE_HEADERS):
        header_cells[col_idx].text = header_text

    # Data rows
    for row_idx, row_data in enumerate(_SAMPLE_TABLE_ROWS):
        row_cells = table.rows[row_idx + 1].cells
        for col_idx, cell_text in enumerate(row_data):
            row_cells[col_idx].text = cell_text

    doc.add_paragraph()   # spacer after table

    # --- Heading 1C (Lists section) ---
    doc.add_paragraph(_SAMPLE_HEADING_1C, style=_STYLE_HEADING_1)

    # --- Heading 2B and bullet list ---
    doc.add_paragraph(_SAMPLE_HEADING_2B, style=_STYLE_HEADING_2)
    for item in _SAMPLE_BULLET_ITEMS:
        doc.add_paragraph(item, style=_STYLE_BULLET)

    # --- Heading 2C and numbered list ---
    doc.add_paragraph(_SAMPLE_HEADING_2C, style=_STYLE_HEADING_2)
    for item in _SAMPLE_NUMBERED_ITEMS:
        doc.add_paragraph(item, style=_STYLE_NUMBER)

    _add_section_break(doc)   # Section 3 ends here -- next page starts references

    # --- References section (Section 4) ---
    # REFERENCES heading + injection marker. formatter.py inserts reference
    # content at [[REFERENCES_START]]. The heading uses Title style so the
    # body scan in writer.py recognises it as the section boundary.
    # REF_ID: RM_DO178_001
    doc.add_paragraph(_REFS_HEADING_TEXT, style=_STYLE_TITLE)
    doc.add_paragraph(_MARKER_REFS_START, style=_STYLE_NORMAL)
    doc.add_paragraph(_MARKER_REFS_NOTE,  style=_STYLE_NORMAL)

    return doc


# ---------------------------------------------------------------------------
# Output path resolver
# Collision-avoids if Rhema_Format_Sample.docx already exists.
# REF_ID: RM_DO178_001
# ---------------------------------------------------------------------------

def _resolve_sample_path() -> Path:
    """Return a collision-free output path for the sample document."""
    DEFAULT_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    candidate: Path = DEFAULT_OUTPUT_DIR / (SAMPLE_OUTPUT_FILENAME + OUTPUT_EXTENSION)
    counter: int = 1
    while candidate.exists() is True:
        candidate = DEFAULT_OUTPUT_DIR / (
            SAMPLE_OUTPUT_FILENAME + "_" + str(counter) + OUTPUT_EXTENSION
        )
        counter += 1
    return candidate


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------

def _prompt(message: str) -> str:
    """Print a prompt and return stripped input. Handles Ctrl+C."""
    print(message)
    try:
        return input("> ").strip()
    except (KeyboardInterrupt, EOFError):
        print()
        return ""


def _offer_save_to_library(sample_path: Path) -> None:
    """
    Ask the user whether to add the generated sample to the template library.
    Routes to the existing save_template flow on confirmation.
    User may skip by pressing Enter or typing anything other than a confirm signal.
    REF_ID: RM_NIST_001 -- no file written to library without explicit confirmation
    """
    raw: str = _prompt(MSG_SAMPLE_SAVE_PROMPT).lower()

    if raw not in PREVIEW_VALID_CONFIRM_SIGNALS:
        print(MSG_SAMPLE_SKIPPED)
        return

    # Name prompt -- loop until non-empty name provided or user cancels
    while True:
        name_raw: str = _prompt(MSG_SAVE_TEMPLATE_NAME_PROMPT)
        if len(name_raw) == 0:
            print(MSG_SAMPLE_SKIPPED)
            return
        break

    ok: bool = save_template(sample_path, name_raw)
    if ok is True:
        print(msg_template_saved(name_raw))
    else:
        print(fmt(LABEL_ERROR,
            "Rhema could not save the sample to the library. "
            "You can add it later via Manage Templates."))


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------

def run_sample_mode() -> None:
    """
    Run the format template generator mode.
    Generates a sample .docx, writes it to the output folder, then offers
    to add it to the template library. Called by main.py on option 4.
    REF_ID: RM_DO178_001 -- defined entry, defined exit, no inference
    """
    print(MSG_SAMPLE_INTRO)
    print(MSG_SAMPLE_GENERATING)

    doc = _build_sample_doc()
    if doc is None:
        print(fmt(LABEL_ERROR,
            "python-docx is not installed. Run: pip install python-docx"))
        return

    output_path: Path = _resolve_sample_path()

    try:
        doc.save(str(output_path))
    except Exception:
        print(MSG_SAMPLE_WRITE_ERROR)
        return

    if output_path.exists() is False or output_path.stat().st_size == 0:
        print(MSG_SAMPLE_WRITE_ERROR)
        return

    print(msg_sample_done(str(output_path)))

    # HITL gate -- offer library save, user may skip. REF_ID: RM_NIST_001
    _offer_save_to_library(output_path)
